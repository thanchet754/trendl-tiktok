"""
Professional Excel Exporter for TikTok Shop US Trends & Ideas
Generates styled .xlsx spreadsheets with multiple sheets, color coding, and clickable links.
Includes Top Videos GMV 24h & Top Influencers sheets.
"""

import os
from datetime import datetime
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def export_to_excel(analyzed_data: Dict[str, Any], output_path: str = None) -> str:
    if not output_path:
        os.makedirs("exports", exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = os.path.join("exports", f"TikTok_Shop_US_Trends_{timestamp}.xlsx")

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1E293B")
    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)
    link_font = Font(name="Calibri", size=10, color="2563EB", underline="single")

    viral_header_fill = PatternFill(start_color="E11D48", end_color="E11D48", fill_type="solid")     # Rose 600
    evergreen_header_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid") # Emerald 600
    all_header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")       # Blue 500
    purple_header_fill = PatternFill(start_color="7C3AED", end_color="7C3AED", fill_type="solid")    # Purple 600
    amber_header_fill = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")     # Amber 600

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    sheets_to_format = []

    # ================= 1. SHEET: TOP 24H VIRAL SPIKES =================
    ws_viral = wb.create_sheet(title="🔥 Top 24h Viral Spikes")
    ws_viral.views.sheetView[0].showGridLines = True
    sheets_to_format.append(ws_viral)

    headers_viral = [
        "STT", "Tên Sản Phẩm / Ý Tưởng", "Ngách Hàng", "Viral Score", "Evergreen", "Giá Tham Khảo", 
        "Nền Tảng Đối Soát", "Góc Quay Hook 3 Giây", "Đối Tượng Khách Hàng (Persona)", "Link Nguồn"
    ]
    
    # Title row
    ws_viral.append(["TOP SẢN PHẨM BÙNG NỔ VIRAL 24H CHO TIKTOK SHOP US"])
    ws_viral.cell(1, 1).font = title_font
    ws_viral.row_dimensions[1].height = 28
    
    # Header row
    ws_viral.append(headers_viral)
    ws_viral.row_dimensions[2].height = 26
    for col_num in range(1, len(headers_viral) + 1):
        cell = ws_viral.cell(2, col_num)
        cell.font = header_font
        cell.fill = viral_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 3
    for i, it in enumerate(analyzed_data.get("viral_24h", []), 1):
        strat = it.get("strategy", {})
        row = [
            i,
            it.get("title", ""),
            it.get("category", "General"),
            it.get("viral_score", 0),
            it.get("evergreen_score", 0),
            it.get("price", it.get("clean_price", "$19.99")),
            ", ".join(it.get("verified_platforms", [])),
            strat.get("hook_angle", it.get("hook_style", "")),
            strat.get("audience", ""),
            it.get("url", "")
        ]
        ws_viral.append(row)
        ws_viral.row_dimensions[row_idx].height = 22
        
        for col_num in range(1, len(row) + 1):
            cell = ws_viral.cell(row_idx, col_num)
            cell.font = normal_font
            cell.border = thin_border
            if col_num in [1, 4, 5, 6]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_num == 10 and cell.value and cell.value.startswith("http"):
                cell.font = link_font
                cell.hyperlink = cell.value
            else:
                cell.alignment = Alignment(vertical="center")
        row_idx += 1

    # ================= 2. SHEET: EVERGREEN WINNERS =================
    ws_ever = wb.create_sheet(title="🌲 Evergreen Winners")
    ws_ever.views.sheetView[0].showGridLines = True
    sheets_to_format.append(ws_ever)

    headers_ever = [
        "STT", "Tên Sản Phẩm / Ý Tưởng", "Ngách Hàng", "Evergreen Score", "Viral Score", "Biên Lợi Nhuận Ước Tính", 
        "Định Dạng Video Khuyên Dùng", "Kêu Gọi Hành Động (CTA)", "Link Nguồn"
    ]
    
    ws_ever.append(["DANH SÁCH SẢN PHẨM BỀN VỮNG (EVERGREEN) CHẠY ADS QUANH NĂM"])
    ws_ever.cell(1, 1).font = title_font
    ws_ever.row_dimensions[1].height = 28

    ws_ever.append(headers_ever)
    ws_ever.row_dimensions[2].height = 26
    for col_num in range(1, len(headers_ever) + 1):
        cell = ws_ever.cell(2, col_num)
        cell.font = header_font
        cell.fill = evergreen_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 3
    for i, it in enumerate(analyzed_data.get("evergreen", []), 1):
        strat = it.get("strategy", {})
        row = [
            i,
            it.get("title", ""),
            it.get("category", "General"),
            it.get("evergreen_score", 0),
            it.get("viral_score", 0),
            strat.get("est_margin", "70% - 85%"),
            strat.get("format_type", "Problem - Agitate - Solution"),
            strat.get("call_to_action", "Click link bio ngay"),
            it.get("url", "")
        ]
        ws_ever.append(row)
        ws_ever.row_dimensions[row_idx].height = 22
        
        for col_num in range(1, len(row) + 1):
            cell = ws_ever.cell(row_idx, col_num)
            cell.font = normal_font
            cell.border = thin_border
            if col_num in [1, 4, 5, 6]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_num == 9 and cell.value and cell.value.startswith("http"):
                cell.font = link_font
                cell.hyperlink = cell.value
            else:
                cell.alignment = Alignment(vertical="center")
        row_idx += 1

    # ================= 3. SHEET: TẤT CẢ Ý TƯỞNG & ĐỐI SOÁT =================
    ws_all = wb.create_sheet(title="📊 Tất Cả Ý Tưởng")
    ws_all.views.sheetView[0].showGridLines = True
    sheets_to_format.append(ws_all)

    headers_all = [
        "STT", "Tên Sản Phẩm", "Phân Loại Trend", "Ngành Hàng", "Tổng Điểm Cơ Hội", "Viral Score", "Evergreen", 
        "Giá Bán", "Nguồn Dữ Liệu Gốc", "Nền Tảng Đối Soát", "Định Dạng Video", "Link Nguồn"
    ]

    ws_all.append(["BẢNG TỔNG HỢP TOÀN BỘ Ý TƯỞNG & ĐỐI SOÁT ĐA NỀN TẢNG"])
    ws_all.cell(1, 1).font = title_font
    ws_all.row_dimensions[1].height = 28

    ws_all.append(headers_all)
    ws_all.row_dimensions[2].height = 26
    for col_num in range(1, len(headers_all) + 1):
        cell = ws_all.cell(2, col_num)
        cell.font = header_font
        cell.fill = all_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 3
    for i, it in enumerate(analyzed_data.get("all_ideas", []), 1):
        strat = it.get("strategy", {})
        row = [
            i,
            it.get("title", ""),
            it.get("label", ""),
            it.get("category", "General"),
            it.get("opportunity_score", 0),
            it.get("viral_score", 0),
            it.get("evergreen_score", 0),
            it.get("price", it.get("clean_price", "$19.99")),
            it.get("source", ""),
            ", ".join(it.get("verified_platforms", [])),
            strat.get("format_type", ""),
            it.get("url", "")
        ]
        ws_all.append(row)
        ws_all.row_dimensions[row_idx].height = 22
        
        for col_num in range(1, len(row) + 1):
            cell = ws_all.cell(row_idx, col_num)
            cell.font = normal_font
            cell.border = thin_border
            if col_num in [1, 5, 6, 7, 8]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_num == 12 and cell.value and cell.value.startswith("http"):
                cell.font = link_font
                cell.hyperlink = cell.value
            else:
                cell.alignment = Alignment(vertical="center")
        row_idx += 1

    # ================= 4. SHEET: TOP VIDEOS GMV 24H =================
    top_videos = analyzed_data.get("top_videos", [])
    if top_videos:
        ws_vids = wb.create_sheet(title="🎬 Top Videos GMV 24h")
        ws_vids.views.sheetView[0].showGridLines = True
        sheets_to_format.append(ws_vids)

        headers_vids = [
            "Rank", "Sản Phẩm Đính Kèm", "Ngành Hàng", "Sub-Niche", "KOC / Creator", "TikTok Handle", 
            "Lượt Xem", "Thời Lượng", "Giá Bán", "Đã Bán 24h", "GMV 24h", "Link Video Trực Tiếp"
        ]

        ws_vids.append(["TOP VIDEOS CÓ DOANH SỐ GMV CAO NHẤT 24H TRÊN TIKTOK SHOP US"])
        ws_vids.cell(1, 1).font = title_font
        ws_vids.row_dimensions[1].height = 28

        ws_vids.append(headers_vids)
        ws_vids.row_dimensions[2].height = 26
        for col_num in range(1, len(headers_vids) + 1):
            cell = ws_vids.cell(2, col_num)
            cell.font = header_font
            cell.fill = purple_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        row_idx = 3
        for v in top_videos:
            row = [
                v.get("rank", 0),
                v.get("product_name", ""),
                v.get("category", ""),
                v.get("sub_niche", ""),
                v.get("creator_name", ""),
                v.get("creator_handle", ""),
                v.get("views", ""),
                v.get("duration", ""),
                v.get("product_price", ""),
                v.get("items_sold_24h", 0),
                v.get("gmv_24h", ""),
                v.get("video_url", "")
            ]
            ws_vids.append(row)
            ws_vids.row_dimensions[row_idx].height = 22

            for col_num in range(1, len(row) + 1):
                cell = ws_vids.cell(row_idx, col_num)
                cell.font = normal_font
                cell.border = thin_border
                if col_num in [1, 7, 8, 9, 10, 11]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_num == 12 and cell.value and cell.value.startswith("http"):
                    cell.font = link_font
                    cell.hyperlink = cell.value
                else:
                    cell.alignment = Alignment(vertical="center")
            row_idx += 1

    # ================= 5. SHEET: TOP INFLUENCERS 24H =================
    top_infs = analyzed_data.get("top_influencers", [])
    if top_infs:
        ws_infs = wb.create_sheet(title="👑 Top Influencers 24h")
        ws_infs.views.sheetView[0].showGridLines = True
        sheets_to_format.append(ws_infs)

        headers_infs = [
            "Rank", "Tên Creator / KOC", "TikTok Handle", "Followers", "Ngành Hàng", "Sub-Niche",
            "Sản Phẩm Top 1", "Đã Bán 24h", "GMV 24h", "Tương Tác", "Link Kênh TikTok"
        ]

        ws_infs.append(["TOP INFLUENCERS / AFFILIATE CREATORS CÓ SỐ BÁN & GMV CAO NHẤT 24H"])
        ws_infs.cell(1, 1).font = title_font
        ws_infs.row_dimensions[1].height = 28

        ws_infs.append(headers_infs)
        ws_infs.row_dimensions[2].height = 26
        for col_num in range(1, len(headers_infs) + 1):
            cell = ws_infs.cell(2, col_num)
            cell.font = header_font
            cell.fill = amber_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        row_idx = 3
        for inf in top_infs:
            row = [
                inf.get("rank", 0),
                inf.get("name", ""),
                inf.get("handle", ""),
                inf.get("followers", ""),
                inf.get("category", ""),
                inf.get("sub_niche", ""),
                inf.get("best_product_title", ""),
                inf.get("items_sold_24h", 0),
                inf.get("gmv_24h", ""),
                inf.get("engagement_rate", ""),
                inf.get("profile_url", "")
            ]
            ws_infs.append(row)
            ws_infs.row_dimensions[row_idx].height = 22

            for col_num in range(1, len(row) + 1):
                cell = ws_infs.cell(row_idx, col_num)
                cell.font = normal_font
                cell.border = thin_border
                if col_num in [1, 4, 8, 9, 10]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_num == 11 and cell.value and cell.value.startswith("http"):
                    cell.font = link_font
                    cell.hyperlink = cell.value
                else:
                    cell.alignment = Alignment(vertical="center")
            row_idx += 1

    # Auto-adjust column widths
    for sheet in sheets_to_format:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col[1:]: # skip title row
                val = str(cell.value or '')
                if cell.hyperlink:
                    val = "Xem Link"
                max_len = max(max_len, len(val))
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

    wb.save(output_path)
    return output_path
